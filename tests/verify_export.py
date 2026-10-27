import openpyxl, sys
sys.path.insert(0,'.')
exp = ['ragione_sociale']  # placeholder
for fn in ['tests/export_full.xlsx','tests/export_oggi.xlsx']:
    wb = openpyxl.load_workbook(fn)
    print('==', fn, wb.sheetnames)
    ws = wb['Lead']
    hdr = [c.value for c in ws[1]]
    print('colonne', len(hdr), hdr)
    print('righe dati', ws.max_row-1)
    ci = {h:i for i,h in enumerate(hdr)}
    hits = []
    fills = {}
    for row in ws.iter_rows(min_row=2):
        col = row[ci['colore']].value
        if col:
            f = row[0].fill.fgColor.rgb if row[0].fill and row[0].fill.fill_type else None
            fl = row[-1].fill.fgColor.rgb if row[-1].fill and row[-1].fill.fill_type else None
            fills[col] = (f, fl)
        if row[ci['referente']].value == 'Sig.ra Bianchi':
            hits.append({h: row[i].value for h,i in ci.items() if h in ('ragione_sociale','citta','visitato','data_visita','esito','referente','note_visita','prossimo_passo','colore','ultima_modifica','modificato_da')})
    print('record modificato:', hits)
    print('fill per colore (prima/ultima cella):', fills)
    print('nuovo negozio presente:', any(r[ci['ragione_sociale']].value=='Gioielleria Test Savona' for r in ws.iter_rows(min_row=2)))
    for sh in ('Riepilogo','Legenda'):
        w = wb[sh]; print('--', sh, [[c.value for c in r] for r in w.iter_rows(max_row=12)])
